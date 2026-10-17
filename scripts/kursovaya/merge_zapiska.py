#!/usr/bin/env python3
"""Сборка общей записки: черновик (разделы 1.1 – 1.3) + дополнение к 1.3 + разделы 1.4 и 1.5.

Использование: python3 merge_zapiska.py черновик.docx выход.docx
Файлы разделов берутся из docs/kursovaya (razdel-1.3-dopolnenie.docx, razdel-1.4.docx, razdel-1.5.docx).
Широкий рисунок (больше 18 см, диаграмма вариантов использования) выводится на отдельной альбомной странице.
Оформление новых разделов приводится к оформлению черновика (стили, шрифт 14, одинарный интервал, таблицы), затем
таблицы разрезаются по страницам с подписью «Продолжение таблицы X.X» (split_tables.py).
"""
import copy, io, os, re, sys, tempfile
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import split_tables

ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
KURS = os.path.join(ROOT, 'docs', 'kursovaya')
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def text_of(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))


def find_par(body, start):
    for el in body:
        if el.tag == qn('w:p') and text_of(el).strip().startswith(start):
            return el
    raise SystemExit(f'не найден абзац: {start}')


def replace_in_runs(doc, old, new):
    for p in doc.paragraphs:
        for r in p.runs:
            if old in r.text and r._r.find(qn('w:drawing')) is None:
                r.text = r.text.replace(old, new)


class Maker:
    """Создаёт элементы в оформлении черновика."""

    def __init__(self, draft):
        P = draft.paragraphs
        self.draft = draft
        self.body_ppr = copy.deepcopy(next(p for p in P if p.text.startswith('Требования к программному средству сформированы'))._p.pPr)
        cap = next(p for p in P if p.text.startswith('Рисунок 1.1 '))
        self.fig_ppr = copy.deepcopy(cap._p.pPr)
        tcap = next(p for p in P if p.text.startswith('Таблица 1.5 '))
        self.tcap_ppr = copy.deepcopy(tcap._p.pPr)
        self.h2 = next(p for p in P if p.style.name == 'Heading 2')._p.pPr.find(qn('w:pStyle')).get(qn('w:val'))
        self.table_style = draft.tables[2]._tbl.tblPr.find(qn('w:tblStyle')).get(qn('w:val'))
        self.docpr = 5000
        self.base_sect = copy.deepcopy(draft.element.body.find(qn('w:sectPr')))

    def par(self, ppr=None, text=''):
        p = OxmlElement('w:p')
        if ppr is not None:
            p.append(copy.deepcopy(ppr))
        if text:
            r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = text
            t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); r.append(t); p.append(r)
        return p

    def sect_par(self, landscape):
        """Абзац-разделитель разделов: sectPr описывает раздел, который заканчивается этим абзацем."""
        sp = copy.deepcopy(self.base_sect)
        if landscape:
            for el in sp.findall(qn('w:titlePg')):
                sp.remove(el)
            pg = sp.find(qn('w:pgSz')); pg.set(qn('w:w'), '16838'); pg.set(qn('w:h'), '11906'); pg.set(qn('w:orient'), 'landscape')
            mar = sp.find(qn('w:pgMar')); mar.set(qn('w:left'), '1134'); mar.set(qn('w:right'), '851'); mar.set(qn('w:top'), '850'); mar.set(qn('w:bottom'), '850')
        ppr = OxmlElement('w:pPr'); ppr.append(sp)
        p = OxmlElement('w:p'); p.append(ppr)
        return p

    def heading(self, text):
        ppr = OxmlElement('w:pPr')
        st = OxmlElement('w:pStyle'); st.set(qn('w:val'), self.h2); ppr.append(st)
        return self.par(ppr, text)

    def picture(self, src_p, src_doc):
        ppr = OxmlElement('w:pPr')
        ppr.append(OxmlElement('w:keepNext'))
        ppr.append(OxmlElement('w:contextualSpacing'))
        jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); ppr.append(jc)
        p = OxmlElement('w:p'); p.append(ppr)
        for r in src_p.findall(qn('w:r')):
            if r.find(qn('w:drawing')) is not None:
                nr = copy.deepcopy(r)
                for blip in nr.iter(f'{{{A}}}blip'):
                    rid = blip.get(f'{{{R}}}embed')
                    blob = src_doc.part.related_parts[rid].blob
                    new_rid, _ = self.draft.part.get_or_add_image(io.BytesIO(blob))
                    blip.set(f'{{{R}}}embed', new_rid)
                for dp in nr.iter('{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr'):
                    self.docpr += 1; dp.set('id', str(self.docpr))
                p.append(nr)
        return p

    def table(self, src_tbl):
        t = copy.deepcopy(src_tbl)
        tblPr = t.find(qn('w:tblPr'))
        for el in tblPr.findall(qn('w:tblStyle')) + tblPr.findall(qn('w:tblW')) + tblPr.findall(qn('w:tblLayout')):
            tblPr.remove(el)
        st = OxmlElement('w:tblStyle'); st.set(qn('w:val'), self.table_style); tblPr.insert(0, st)
        w = OxmlElement('w:tblW'); w.set(qn('w:w'), '9350'); w.set(qn('w:type'), 'dxa'); tblPr.append(w)
        lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
        rows = t.findall(qn('w:tr'))
        # ширины колонок: по ширине ячеек первой строки, приведённые к 9350
        tcw = [int(tc.find(qn('w:tcPr')).find(qn('w:tcW')).get(qn('w:w'))) for tc in rows[0].findall(qn('w:tc'))]
        total = sum(tcw); widths = [round(x * 9350 / total) for x in tcw]; widths[-1] += 9350 - sum(widths)
        grid = t.find(qn('w:tblGrid'))
        for gc, wd in zip(grid.findall(qn('w:gridCol')), widths):
            gc.set(qn('w:w'), str(wd))
        for tr in rows:
            for tc, wd in zip(tr.findall(qn('w:tc')), widths):
                tcw_el = tc.find(qn('w:tcPr')).find(qn('w:tcW')); tcw_el.set(qn('w:w'), str(wd)); tcw_el.set(qn('w:type'), 'dxa')
        for i, tr in enumerate(rows):
            for p in tr.iter(qn('w:p')):
                ppr = p.find(qn('w:pPr'))
                if ppr is None:
                    ppr = OxmlElement('w:pPr'); p.insert(0, ppr)
                for el in list(ppr):
                    if el.tag in (qn('w:spacing'), qn('w:jc'), qn('w:contextualSpacing')):
                        ppr.remove(el)
                sp = OxmlElement('w:spacing'); sp.set(qn('w:line'), '240'); sp.set(qn('w:lineRule'), 'auto'); ppr.append(sp)
                ppr.append(OxmlElement('w:contextualSpacing'))
                jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center' if i == 0 else 'left'); ppr.append(jc)
                for r in p.findall(qn('w:r')):
                    rpr = r.find(qn('w:rPr'))
                    if rpr is None:
                        rpr = OxmlElement('w:rPr'); r.insert(0, rpr)
                    for el in list(rpr):
                        if el.tag in (qn('w:sz'), qn('w:szCs'), qn('w:b'), qn('w:rFonts')):
                            rpr.remove(el)
                    for name in ('w:sz', 'w:szCs'):
                        z = OxmlElement(name); z.set(qn('w:val'), '28'); rpr.append(z)
        return t

    def convert(self, path, scales=None):
        """Приводит элементы раздела к оформлению черновика."""
        flat = tempfile.mktemp(suffix='.docx')
        split_tables.prepare(path, flat)
        src = docx.Document(flat)
        out = []
        for el in src.element.body:
            if el.tag == qn('w:tbl'):
                out.append(self.table(el))
            elif el.tag == qn('w:p'):
                text = text_of(el).strip()
                has_pic = el.find('.//' + qn('w:drawing')) is not None
                if has_pic:
                    out.append(self.picture(el, src))
                elif re.match(r'^1\.\d+(\.\d+)?\s', text):
                    num, rest = text.split(' ', 1)
                    out.append(self.heading(f'{num}\t{rest}'))
                elif text.startswith('Рисунок '):
                    out.append(self.par(self.fig_ppr, text))
                elif text.startswith('Таблица '):
                    out.append(self.par(self.tcap_ppr, text))
                elif not text:
                    out.append(self.par(None))
                else:
                    out.append(self.par(self.body_ppr, text))
        os.remove(flat)
        apply_scales(out, scales or {})
        return out


EMU_CM = 360000
NS_WP = '{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}'


def figure_number(els, i):
    """Номер рисунка по подписи, которая идёт следом за абзацем с картинкой."""
    for e in els[i + 1:i + 4]:
        m = re.match(r'\s*Рисунок 1\.(\d+)\s', text_of(e))
        if m:
            return int(m.group(1))
    return None


def picture_height_cm(el):
    return int(next(el.iter(NS_WP + 'extent')).get('cy')) / EMU_CM


def apply_scales(els, scales):
    """Уменьшает рисунок (множитель <= 1), чтобы он поместился на странице."""
    for i, e in enumerate(els):
        if e.tag == qn('w:p') and e.find('.//' + qn('w:drawing')) is not None:
            n = figure_number(els, i)
            k = scales.get(n)
            if not k:
                continue
            for ext in e.iter(NS_WP + 'extent'):
                ext.set('cx', str(round(int(ext.get('cx')) * k))); ext.set('cy', str(round(int(ext.get('cy')) * k)))
            for ext in e.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}ext'):
                if ext.get('cx'):
                    ext.set('cx', str(round(int(ext.get('cx')) * k))); ext.set('cy', str(round(int(ext.get('cy')) * k)))


used_landscape = [False]


def normalize(mk, els, moves=None):
    """Пустые строки как в черновике: вокруг заголовков, рисунков и таблиц."""
    def kind(el):
        if el.tag == qn('w:tbl'):
            return 'tbl'
        t = text_of(el).strip()
        ppr0 = el.find(qn('w:pPr'))
        if ppr0 is not None and ppr0.find(qn('w:sectPr')) is not None:
            return 'sect'
        if el.find('.//' + qn('w:drawing')) is not None:
            return 'pic'
        if not t:
            return 'blank'
        ppr = el.find(qn('w:pPr'))
        if ppr is not None and ppr.find(qn('w:pStyle')) is not None and ppr.find(qn('w:pStyle')).get(qn('w:val')) == mk.h2:
            return 'h'
        if t.startswith('Рисунок '):
            return 'figcap'
        if t.startswith('Таблица '):
            return 'tcap'
        return 'p'
    def wide(el):
        return any(int(x.get('cx')) > 6480000 for x in el.iter('{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent'))
    items = [(kind(e), e) for e in els if kind(e) != 'blank']
    for n, k in (moves or {}).items():      # рисунок n сдвигается ниже на k абзацев или таблиц
        cap = next((i for i, (kk, e) in enumerate(items) if kk == 'figcap' and text_of(e).strip().startswith(f'Рисунок 1.{n} ')), None)
        if cap is None or items[cap - 1][0] != 'pic':
            continue
        block = items[cap - 1:cap + 1]
        rest = items[cap + 1:]
        del items[cap - 1:]
        pos = 0
        for _ in range(k):
            if pos < len(rest) and rest[pos][0] == 'p':
                pos += 1
            elif pos + 1 < len(rest) and rest[pos][0] == 'tcap' and rest[pos + 1][0] == 'tbl':
                pos += 2
            else:
                break
        items += rest[:pos] + block + rest[pos:]
    out = []
    pending_landscape = False
    used_landscape[0] = False
    def blank():
        if out and kind(out[-1]) != 'blank':
            out.append(mk.par(None))
    for i, (k, e) in enumerate(items):
        if k == 'pic' and wide(e):
            out.append(mk.sect_par(False))      # всё до рисунка – книжный раздел
            pending_landscape = True
            used_landscape[0] = True
        elif k in ('h', 'pic', 'tcap'):
            blank()
        out.append(e)
        if k in ('h', 'tbl', 'figcap'):
            out.append(mk.par(None))
        if k == 'pic':
            out.append(mk.par(None))
        if k == 'figcap' and pending_landscape:
            out.append(mk.sect_par(True))       # рисунок и подпись – альбомный раздел
            pending_landscape = False
    return out


def assemble_raw(draft_path, scales=None, moves=None):
    """Собирает общий документ без разрезания таблиц; возвращает путь к временному файлу."""
    scales = scales or {}
    moves = moves or {}
    draft = docx.Document(draft_path)
    # таблица функциональных требований становится 1.7 (перед ней вставляются таблицы 1.5 и 1.6)
    mk = Maker(draft)
    replace_in_runs(draft, 'таблице 1.5', 'таблице 1.7')
    replace_in_runs(draft, 'Таблица 1.5 –', 'Таблица 1.7 –')
    body = draft.element.body

    # диаграммы вариантов использования из 1.3 убираются: одна диаграмма приведена в 1.5
    first = find_par(body, 'Функциональность программного средства для разных категорий пользователей')
    last = find_par(body, 'Спецификация функциональных требований приведена')
    el = first
    while el is not last:
        nxt = el.getnext()
        body.remove(el)
        el = nxt
    import build_15
    uc_no = build_15.FIRST + build_15.layout()['order'].index('uc')
    last.addprevious(mk.par(mk.body_ppr, f'Функциональность программного средства для разных категорий пользователей представлена диаграммой вариантов использования, приведённой в подразделе 1.5 (рисунок 1.{uc_no}).'))

    # дополнение к 1.3: перед ним стоит жёлтая служебная пометка (она не переносится)
    flat = tempfile.mktemp(suffix='.docx')
    split_tables.prepare(os.path.join(KURS, 'razdel-1.3-dopolnenie.docx'), flat)
    sup = docx.Document(flat)
    parts, cur = [], None
    for el in sup.element.body:
        if el.tag == qn('w:sectPr'):
            continue
        if el.tag == qn('w:p') and el.find('.//' + qn('w:highlight')) is not None:
            cur = []; parts.append(cur); continue
        if cur is not None:
            cur.append(copy.deepcopy(el))
    os.remove(flat)
    assert len(parts) == 1, len(parts)

    anchor_a = find_par(body, 'Программное средство должно быть реализовано как клиент-серверное')
    end_heading = find_par(body, 'СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ')
    anchor_end = end_heading.getprevious()           # абзац с разрывом страницы
    parts = [normalize(mk, part) for part in parts]
    for el in reversed(parts[0]):
        anchor_a.addnext(el)
    new = normalize(mk, mk.convert(os.path.join(KURS, 'razdel-1.4.docx'), scales) + mk.convert(os.path.join(KURS, 'razdel-1.5.docx'), scales), moves)
    for el in new:
        anchor_end.addprevious(el)

    if used_landscape[0]:
        # финальный раздел идёт после альбомного: номер страницы должен быть и на его первой странице
        for el in body.find(qn('w:sectPr')).findall(qn('w:titlePg')):
            body.find(qn('w:sectPr')).remove(el)
    raw = tempfile.mktemp(suffix='.docx')
    draft.save(raw)
    return raw


def lo_path(out_path):
    return out_path.replace('.docx', '.lo.docx')


def assemble(draft_path, out_path, scales, moves):
    """out_path – docx для Word; рядом пишется копия .lo.docx для LibreOffice (измерение вёрстки и PDF)."""
    raw = assemble_raw(draft_path, scales, moves)
    split_tables.main(raw, out_path, lo_copy=lo_path(out_path))
    os.remove(raw)


def figure_heights(draft_path):
    """Высоты рисунков разделов 1.4 и 1.5 (см) по номеру."""
    mk_doc = docx.Document(draft_path)
    mk = Maker(mk_doc)
    heights = {}
    for name in ('razdel-1.4.docx', 'razdel-1.5.docx'):
        els = mk.convert(os.path.join(KURS, name))
        for i, e in enumerate(els):
            if e.tag == qn('w:p') and e.find('.//' + qn('w:drawing')) is not None:
                heights[figure_number(els, i)] = picture_height_cm(e)
    return heights


def main(draft_path, out_path):
    """Сборка с подгонкой вёрстки: рисунок, который уходит на следующую страницу и оставляет внизу пустое место,
    уменьшается (не более чем на 15 %), чтобы поместиться. Порядок абзацев и рисунков задан вручную в MOVES."""
    import check_layout
    heights = figure_heights(draft_path)
    scales, seen = {}, set()
    for it in range(6):
        assemble(draft_path, out_path, scales, MOVES)
        gaps = check_layout.analyze(lo_path(out_path))
        todo = []
        for g in gaps:
            m = re.match(r'Рисунок 1\.(\d+)\s', g['next'])
            if g['gap'] > LIMIT and m and int(m.group(1)) >= 10:
                todo.append((int(m.group(1)), g['gap']))
        print('подгонка', it, 'пустые места перед рисунками:', todo, 'масштабы', scales)
        changed = False
        for n, gap in todo:
            need = heights[n] * scales.get(n, 1.0) + 2.6        # рисунок, подпись и пустые строки
            k = (gap - 2.6) / heights[n]
            if need - gap <= 0.15 * heights[n] and k >= 0.85 and k < scales.get(n, 1.0):
                scales[n] = round(k, 3); changed = True
        if not changed or tuple(sorted(scales.items())) in seen:
            break
        seen.add(tuple(sorted(scales.items())))
    pdf = check_layout.to_pdf(lo_path(out_path))
    os.replace(pdf, out_path.replace('.docx', '.pdf'))
    for g in check_layout.analyze(out_path.replace('.docx', '.pdf')):
        if g['gap'] > LIMIT:
            print(f"осталось пустое место: стр. {g['page']}, {g['gap']} см, дальше {g['next'][:50]!r}")


# Ручной порядок: рисунок N сдвигается ниже на k абзацев/таблиц, чтобы текст заполнил страницу
MOVES = {}

LIMIT = 5.0

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
