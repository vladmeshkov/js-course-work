"""Дополнение к разделу 1.3: бизнес-требования и описание вариантов использования.
Строит docx в стиле черновика записки (шаблон: сам черновик), затем split_tables.py разрезает таблицы по страницам
с подписью «Продолжение таблицы X».

Использование: python3 build_13.py шаблон.docx выход.docx
Шаблон – docx черновика записки (из него берутся стили, нумерация списков и параметры страницы).
"""
import copy, os, subprocess, sys, tempfile
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
TPL = sys.argv[1]
OUT = sys.argv[2]
exec(open(os.path.join(HERE, 'content_13.py'), encoding='utf-8').read())   # даёт CONTENT

# ---------------------------------------------------------------- сборка документа
def new_doc():
    d = docx.Document(TPL)
    P = d.paragraphs
    protos = {'body': copy.deepcopy(P[195]._p.pPr), 'cap': copy.deepcopy(P[213]._p.pPr), 'li': copy.deepcopy(P[131]._p.pPr)}
    body = d.element.body
    for el in list(body):
        if not el.tag.endswith('sectPr'):
            body.remove(el)
    # убрать рисунки черновика, чтобы файл не весил 9 МБ
    for rid, rel in list(d.part.rels.items()):
        if 'image' in rel.reltype:
            del d.part.rels[rid]
    return d, protos

def para(d, protos, kind, text='', tag=None):
    p = d.add_paragraph()
    if kind in protos:
        old = p._p.pPr
        if old is not None:
            p._p.remove(old)
        p._p.insert(0, copy.deepcopy(protos[kind]))
    if text:
        p.add_run(text)
    return p

def set_cell(cell, text_lines, align='left', tag=None, keep_next=False):
    cell.text = ''
    first = True
    for line in text_lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        pf = p.paragraph_format
        pPr = p._p.get_or_add_pPr()
        sp = OxmlElement('w:spacing'); sp.set(qn('w:line'), '240'); sp.set(qn('w:lineRule'), 'auto'); pPr.append(sp)
        cs = OxmlElement('w:contextualSpacing'); pPr.append(cs)
        if keep_next:
            kn = OxmlElement('w:keepNext'); pPr.insert(0, kn)
        jc = OxmlElement('w:jc'); jc.set(qn('w:val'), align); pPr.append(jc)
        r = p.add_run(line); r.font.size = Pt(14)
    if tag:
        r = cell.paragraphs[0].add_run(tag); r.font.size = Pt(1); r.font.color.rgb = RGBColor(255, 255, 255)

def make_table(d, header, widths, rows, tagprefix=None, row_offset=0):
    t = d.add_table(rows=1 + len(rows), cols=len(header))
    t.style = d.styles['Table Grid']
    tblPr = t._tbl.tblPr
    for el in tblPr.findall(qn('w:tblW')):
        tblPr.remove(el)
    w = OxmlElement('w:tblW'); w.set(qn('w:w'), str(sum(widths))); w.set(qn('w:type'), 'dxa'); tblPr.append(w)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    grid = t._tbl.tblGrid
    for gc, wd in zip(grid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(wd))
    for i, tr in enumerate(t._tbl.tr_lst):
        trPr = tr.get_or_add_trPr()
        if i == 0:
            trPr.append(OxmlElement('w:cantSplit')); trPr.append(OxmlElement('w:tblHeader'))
        else:
            trPr.append(OxmlElement('w:cantSplit'))
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]; c.width = docx.shared.Twips(widths[j])
        set_cell(c, [h], 'center', keep_next=True)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            c = t.rows[i].cells[j]; c.width = docx.shared.Twips(widths[j])
            lines = val if isinstance(val, list) else [val]
            set_cell(c, lines, 'left', tag=(f'@@{tagprefix}R{row_offset + i - 1}@@' if tagprefix and j == 0 else None), keep_next=(i == 1))
    return t

def build(path):
    d, protos = new_doc()
    for item in CONTENT:
        kind = item[0]
        if kind == 'note':
            p = para(d, protos, 'body')
            r = p.add_run(item[1]); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif kind == 'p':
            para(d, protos, 'body', item[1])
        elif kind == 'li':
            para(d, protos, 'li', item[1])
        elif kind == 'table':
            _, tid, number, title, header, widths, rows = item
            para(d, protos, 'cap').add_run(f'Таблица {number} – {title}')
            make_table(d, header, widths, rows)
            para(d, protos, 'body')    # пустая строка после таблицы
    d.save(path)

def main():
    tmp = tempfile.mkdtemp(prefix='build13_')
    raw = os.path.join(tmp, 'raw.docx')
    build(raw)
    sys.path.insert(0, HERE)
    from common import polish
    polish(raw)
    sys.argv = ['split_tables.py']
    sys.path.insert(0, HERE)
    import split_tables
    split_tables.main(raw, OUT)


main()
