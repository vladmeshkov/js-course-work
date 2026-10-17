#!/usr/bin/env python3
"""Оформление таблиц по правилам записки: если таблица переходит на новую страницу, перед продолжением
пишется «Продолжение таблицы X.X», а шапка таблицы повторяется.

Word не умеет вставлять такую подпись сам, поэтому таблица режется на отдельные таблицы по границам страниц.
Границы страниц берутся из PDF, который строит LibreOffice (в Word разбивка может отличаться на строку-две,
после правок текста запустите скрипт заново на исходном файле).

Использование: python3 split_tables.py вход.docx выход.docx
Заголовок таблицы определяется абзацем «Таблица X – …» непосредственно перед таблицей, шапкой считается первая строка.
"""
import copy, os, re, shutil, subprocess, sys, tempfile
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def caption_of(tbl):
    prev = tbl.getprevious()
    if prev is None or prev.tag != qn('w:p'):
        return None
    text = ''.join(t.text or '' for t in prev.iter(qn('w:t')))
    m = re.match(r'\s*Таблица\s+([\dА-Я.]+)\s*[–-]', text)
    return (prev, m.group(1)) if m else None


def make_continuation(cap_p, number):
    p = copy.deepcopy(cap_p)
    runs = p.findall(qn('w:r'))
    for r in runs[1:]:
        p.remove(r)
    for t in runs[0].findall(qn('w:t')):
        runs[0].remove(t)
    t = OxmlElement('w:t'); t.text = f'Продолжение таблицы {number}'; t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    runs[0].append(t)
    pPr = p.find(qn('w:pPr'))
    if pPr is not None:
        for ind in pPr.findall(qn('w:ind')):
            ind.set(qn('w:left'), '0')
            for a in ('w:hanging', 'w:firstLine'):
                if ind.get(qn(a)) is not None:
                    del ind.attrib[qn(a)]
        for jc in pPr.findall(qn('w:jc')):
            jc.set(qn('w:val'), 'left')
        if pPr.find(qn('w:keepNext')) is None:
            pPr.insert(0, OxmlElement('w:keepNext'))
    return p


def set_row_flags(tr, header):
    trPr = tr.find(qn('w:trPr'))
    if trPr is None:
        trPr = OxmlElement('w:trPr'); tr.insert(0, trPr)
    for tag in ('w:cantSplit', 'w:tblHeader'):
        for el in trPr.findall(qn(tag)):
            trPr.remove(el)
    trPr.append(OxmlElement('w:cantSplit'))
    if header:
        trPr.append(OxmlElement('w:tblHeader'))


def keep_with_next(tr):
    for p in tr.iter(qn('w:p')):
        pPr = p.find(qn('w:pPr'))
        if pPr is None:
            pPr = OxmlElement('w:pPr'); p.insert(0, pPr)
        if pPr.find(qn('w:keepNext')) is None:
            pPr.insert(0, OxmlElement('w:keepNext'))


def tag_row(tr, tag):
    p = next(tr.iter(qn('w:p')))
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), 'FFFFFF'); rPr.append(color)
    for name in ('w:sz', 'w:szCs'):
        sz = OxmlElement(name); sz.set(qn('w:val'), '2'); rPr.append(sz)
    r.append(rPr)
    t = OxmlElement('w:t'); t.text = ' ' + tag; t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); r.append(t)
    p.append(r)


def apply(src, dst, splits, tagged, for_lo=False):
    d = docx.Document(src)
    body = d.element.body
    infos = []
    for idx, tbl in enumerate([e for e in body if e.tag == qn('w:tbl')]):
        cap = caption_of(tbl)
        if cap is None:
            continue
        infos.append((idx, tbl, cap))
    for idx, tbl, (cap_p, number) in infos:
        rows = tbl.findall(qn('w:tr'))
        header, data = rows[0], rows[1:]
        starts = splits.get(idx, [0])
        bounds = starts + [len(data)]
        for tr in rows:
            set_row_flags(tr, False)
        set_row_flags(header, True)
        keep_with_next(header)
        if data:
            keep_with_next(data[0])        # в начале каждой части минимум две строки данных
        chunks = [data[bounds[ci]:bounds[ci + 1]] for ci in range(len(starts))]
        anchor = tbl
        for ci in range(1, len(chunks)):
            new_tbl = copy.deepcopy(tbl)
            for tr in new_tbl.findall(qn('w:tr')):
                new_tbl.remove(tr)
            new_tbl.append(copy.deepcopy(header))
            for tr in chunks[ci]:
                new_tbl.append(tr)
            if len(chunks[ci]) > 1:
                keep_with_next(chunks[ci][0])          # строка переносится из исходной таблицы
            cont = make_continuation(cap_p, number)
            anchor.addnext(cont)
            cont.addnext(new_tbl)
            anchor = new_tbl
        if tagged:
            for ci, chunk in enumerate(chunks):
                for k, tr in enumerate(chunk):
                    tag_row(tr, f'@@T{idx:02d}R{bounds[ci] + k:03d}')
    if for_lo:
        # LibreOffice теряет строки таблицы при сочетании «строка не разрывается» и «не отрывать от следующей»:
        # в копии для измерения и PDF абзацы таблиц без «не отрывать от следующей» (в docx для Word она остаётся)
        for tbl in [e for e in body if e.tag == qn('w:tbl')]:
            for kn in list(tbl.iter(qn('w:keepNext'))):
                kn.getparent().remove(kn)
    d.save(dst)
    return [i for i, _, _ in infos], {i: len(list(t.findall(qn('w:tr')))) for i, t, _ in infos}


def render_pages(path, outdir):
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', outdir, path], check=True, capture_output=True)
    pdf = os.path.join(outdir, os.path.basename(path).replace('.docx', '.pdf'))
    html = pdf.replace('.pdf', '.html')
    subprocess.run(['pdftotext', '-bbox', pdf, html], check=True)
    # слова страницы склеиваются: pdftotext иногда режет белую метку на два слова
    texts, page = {}, 0
    for line in open(html, encoding='utf-8'):
        if '<page ' in line:
            page += 1
            texts[page] = []
        m = re.search(r'<word[^>]*>(.*?)</word>', line)
        if m:
            texts[page].append(m.group(1))
    pages = {}
    for pg, words in texts.items():
        for m in re.finditer(r'@@T(\d\d)R(\d\d\d)', ''.join(words)):
            pages[(int(m.group(1)), int(m.group(2)))] = pg
    return pages


LAST_SPLITS = {}      # разбивка таблиц последнего запуска main (индекс таблицы -> номера строк начала частей)


def unsplit_doc(d):
    """Склеивает таблицы, уже разрезанные «Продолжением», обратно (чтобы скрипт можно было запускать повторно)."""
    body = d.element.body
    for tbl in [e for e in body if e.tag == qn('w:tbl')]:
        prev = tbl.getprevious()
        if prev is None or prev.tag != qn('w:p'):
            continue
        text = ''.join(t.text or '' for t in prev.iter(qn('w:t')))
        if not text.strip().startswith('Продолжение таблицы'):
            continue
        before = prev.getprevious()
        while before is not None and before.tag != qn('w:tbl'):
            before = before.getprevious()
        if before is None:
            continue
        for tr in tbl.findall(qn('w:tr'))[1:]:
            before.append(tr)
        body.remove(prev)
        body.remove(tbl)


def prepare(src, dst):
    d = docx.Document(src)
    unsplit_doc(d)
    d.save(dst)


def main(src, dst, lo_copy=None):
    tmp = tempfile.mkdtemp(prefix='split_tables_')
    flat = os.path.join(tmp, 'flat.docx')
    prepare(src, flat)
    src = flat
    work = os.path.join(tmp, 'work.docx')
    ids, _ = apply(src, work, {}, False)
    d = docx.Document(src)
    counts = {}
    for idx, tbl in enumerate([e for e in d.element.body if e.tag == qn('w:tbl')]):
        if caption_of(tbl) is not None:
            counts[idx] = len(tbl.findall(qn('w:tr'))) - 1
    splits, seen = {}, []
    for it in range(25):
        apply(src, work, splits, True, for_lo=True)
        pages = render_pages(work, tmp)
        new = {}
        for idx, n in counts.items():
            pg = [pages.get((idx, r)) for r in range(n)]
            if None in pg:
                raise SystemExit(f'не найдена метка строки таблицы {idx}: {pg}')
            new[idx] = [0] + [k for k in range(1, n) if pg[k] != pg[k - 1]]
        print('итерация', it, {k: v for k, v in new.items() if len(v) > 1})
        if new == splits:
            break
        if new in seen:
            print('разбивка зациклилась, беру текущий вариант')
            splits = new
            break
        seen.append(copy.deepcopy(new))
        splits = new
    apply(src, dst, splits, False)
    if lo_copy:
        apply(src, lo_copy, splits, False, for_lo=True)
    shutil.rmtree(tmp, ignore_errors=True)
    LAST_SPLITS.clear(); LAST_SPLITS.update(splits)
    print('готово:', dst, 'таблиц с подписью:', len(counts), 'разрезано:', sum(1 for v in splits.values() if len(v) > 1))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
